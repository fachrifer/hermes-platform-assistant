import json

import pytest

from core.agent import HermesAgent
from core.db import Database
from core.finance_import import FinanceImportService
from core.telegram_bot import TelegramInterface


@pytest.mark.asyncio
async def test_recommend_pending_scan_uses_existing_amount_without_staging():
    class Import:
        def __init__(self):
            self.saved = None

        async def resolve_category(self, **kwargs):
            return ("Makanan", "Cafe", True)

        def stage_scanned_transaction(self, tx):
            self.saved = tx
            return 17

    agent = HermesAgent.__new__(HermesAgent)
    agent.finance_import = Import()

    recommendation = await agent.recommend_scanned_transaction(
        {"email_id": "mail-1", "amount": 20000, "type": "expense", "category": "Lainnya",
         "subcategory": "-", "merchant": "Famima", "description": "Transfer"},
        "catat sebagai kopi FamilyMart, sesuaikan kategori",
    )

    assert recommendation["amount"] == 20000
    assert recommendation["category"] == "Makanan"
    assert recommendation["subcategory"] == "Cafe"
    assert agent.finance_import.saved is None


def test_pending_scan_selection_matches_user_description():
    items = [
        {"description": "Transfer ke KRISTYAWAN - famima", "merchant": "", "subject": "Transfer"},
        {"description": "Pembayaran QRIS", "merchant": "", "subject": "Internet Transaction Journal"},
    ]

    selected = TelegramInterface.select_pending_transaction(items, "catat sebagai kopi family mart")

    assert selected is items[0]


@pytest.mark.asyncio
async def test_resolve_category_uses_local_recommendations_without_llm(tmp_path):
    from openpyxl import Workbook
    from core.category_recommendations import CategoryRecommendationLookup
    from core.spreadsheet_finance import CategoryReference, SpreadsheetFinanceService

    categories = tmp_path / "categories.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Kategori Pengeluaran", "Subkategori"])
    sheet.append(["Makanan", "Cafe"])
    workbook.save(categories)

    rec_path = tmp_path / "category_recommendations.json"
    rec_path.write_text(json.dumps({
        "phrases": [],
        "keywords": [{
            "match": "kopi",
            "category": "Makanan",
            "subcategory": "Cafe",
            "count": 10,
            "total": 10,
            "confidence": 1.0,
        }],
    }), encoding="utf-8")

    class LLM:
        async def recommend_category(self, **kwargs):
            raise AssertionError("LLM should not be called for local lookup")

    service = FinanceImportService(
        Database(str(tmp_path / "state.db")),
        None,
        LLM(),
        spreadsheet_service=SpreadsheetFinanceService(
            str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
        ),
        category_recommendations=CategoryRecommendationLookup(
            str(rec_path),
            category_reference=CategoryReference(str(categories)),
        ),
    )

    category, subcategory, ok = await service.resolve_category(
        merchant="Warung",
        description="beli kopi",
        amount=23000,
        tx_type="expense",
        category="Cafe",
        subcategory="-",
        force_recommend=True,
        allow_llm=False,
    )

    assert ok is True
    assert category == "Makanan"
    assert subcategory == "Cafe"


@pytest.mark.asyncio
async def test_recommend_categories_for_batch_updates_invalid_rows(tmp_path):
    from openpyxl import Workbook
    from core.category_recommendations import CategoryRecommendationLookup
    from core.spreadsheet_finance import CategoryReference, SpreadsheetFinanceService

    categories = tmp_path / "categories.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Kategori Pengeluaran", "Subkategori"])
    sheet.append(["Makanan", "Restoran"])
    workbook.save(categories)

    rec_path = tmp_path / "category_recommendations.json"
    rec_path.write_text(json.dumps({
        "phrases": [],
        "keywords": [{
            "match": "soto",
            "category": "Makanan",
            "subcategory": "Restoran",
            "count": 5,
            "total": 5,
            "confidence": 1.0,
        }],
    }), encoding="utf-8")

    db = Database(str(tmp_path / "state.db"))
    batch = db.get_or_create_import_batch("2026-07")
    tx_id = db.stage_transaction(batch["id"], {
        "source_email_id": "A:1",
        "account": "Cash",
        "transaction_date": "2026-07-15",
        "amount": 63000,
        "type": "expense",
        "category": "Makanan & Minuman",
        "subcategory": "-",
        "merchant": "Soto bang yos",
        "description": "Soto",
        "needs_review": 1,
    })
    service = FinanceImportService(
        db, None, None,
        spreadsheet_service=SpreadsheetFinanceService(
            str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
        ),
        category_recommendations=CategoryRecommendationLookup(
            str(rec_path),
            category_reference=CategoryReference(str(categories)),
        ),
    )

    result = await service.recommend_categories_for_batch("2026-07")
    row = db.get_staged_transaction(tx_id)

    assert result["updated"] == 1
    assert row["category"] == "Makanan"
    assert row["subcategory"] == "Restoran"
    assert row["needs_review"] == 0


@pytest.mark.asyncio
async def test_convert_transfer_to_expense_uses_living_cost_category(tmp_path):
    from openpyxl import Workbook
    from core.category_recommendations import CategoryRecommendationLookup
    from core.spreadsheet_finance import CategoryReference, SpreadsheetFinanceService

    categories = tmp_path / "categories.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Kategori Pengeluaran", "Subkategori"])
    sheet.append(["Keluarga", "Istri", "Orang Tua"])
    sheet.append(["Penarikan Dana dan Transfer", "Biaya Transfer", "Uang Elektronik"])
    workbook.save(categories)

    rec_path = tmp_path / "category_recommendations.json"
    rec_path.write_text(json.dumps({
        "phrases": [],
        "keywords": [{
            "match": "nurizka",
            "category": "Keluarga",
            "subcategory": "Istri",
            "count": 5,
            "total": 5,
            "confidence": 1.0,
        }],
    }), encoding="utf-8")

    db = Database(str(tmp_path / "state.db"))
    batch = db.get_or_create_import_batch("2025-09")
    tx_id = db.stage_transaction(batch["id"], {
        "source_email_id": "T:1",
        "account": "Cash",
        "transaction_date": "2025-09-02",
        "amount": 4002500,
        "type": "transfer",
        "category": "Tabungan - Istri",
        "subcategory": "-",
        "merchant": "Nurizka Khoerani",
        "description": "Transfer dengan BI Fast ke NURIZKA KHOERANI",
        "needs_review": 0,
    })
    service = FinanceImportService(
        db, None, None,
        spreadsheet_service=SpreadsheetFinanceService(
            str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
        ),
        category_recommendations=CategoryRecommendationLookup(
            str(rec_path),
            category_reference=CategoryReference(str(categories)),
        ),
    )

    result = await service.convert_transfer_to_expense(tx_id)
    row = db.get_staged_transaction(tx_id)

    assert result["type"] == "expense"
    assert row["type"] == "expense"
    assert row["category"] == "Keluarga"
    assert row["subcategory"] == "Istri"
    assert row["merchant"] == "Bulanan Istri"
    assert row["account"] == "Cash"


@pytest.mark.asyncio
async def test_convert_transfer_to_expense_mama_uses_bulanan_mama(tmp_path):
    from openpyxl import Workbook
    from core.category_recommendations import CategoryRecommendationLookup
    from core.spreadsheet_finance import CategoryReference, SpreadsheetFinanceService

    categories = tmp_path / "categories.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Kategori Pengeluaran", "Subkategori"])
    sheet.append(["Keluarga", "Istri", "Orang Tua"])
    workbook.save(categories)

    rec_path = tmp_path / "category_recommendations.json"
    rec_path.write_text(json.dumps({"phrases": [], "keywords": []}), encoding="utf-8")

    db = Database(str(tmp_path / "state.db"))
    batch = db.get_or_create_import_batch("2025-09")
    tx_id = db.stage_transaction(batch["id"], {
        "source_email_id": "T:2",
        "account": "Cash",
        "transaction_date": "2025-09-01",
        "amount": 1200000,
        "type": "transfer",
        "category": "Cash",
        "subcategory": "-",
        "merchant": "",
        "description": "Transfer ke RATNA NINGSIH",
        "needs_review": 0,
    })
    service = FinanceImportService(
        db, None, None,
        spreadsheet_service=SpreadsheetFinanceService(
            str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
        ),
        category_recommendations=CategoryRecommendationLookup(
            str(rec_path),
            category_reference=CategoryReference(str(categories)),
        ),
    )

    result = await service.convert_transfer_to_expense(tx_id)
    row = db.get_staged_transaction(tx_id)
    assert result["category"] == "Keluarga"
    assert result["subcategory"] == "Orang Tua"
    assert row["merchant"] == "Bulanan Mama"


def test_batch_card_shows_category_recommendation():
    bot = TelegramInterface(agent=None, briefing=None)
    text = bot._format_batch_card(
        {
            "id": 12,
            "transaction_date": "2026-07-15",
            "type": "expense",
            "amount": 23000,
            "category": "Makanan",
            "subcategory": "Cafe",
            "merchant": "Kopi",
            "account": "Cash",
            "description": "Latte",
            "needs_review": 0,
            "possible_duplicate": 0,
        },
        index=1,
        total=7,
        period="2026-07",
        validation={
            "lookup_category": "Makanan",
            "lookup_subcategory": "Cafe",
            "category": "Makanan",
            "subcategory": "Cafe",
            "agreed": True,
            "note": "cocok",
        },
    )
    assert "Local lookup: Makanan / Cafe" in text
    assert "LLM validation: agrees" in text
    assert "Note: Kopi" in text
    assert "Description: Latte" in text
    assert "Final decision is yours" in text
