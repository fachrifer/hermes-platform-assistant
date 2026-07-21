from openpyxl import Workbook, load_workbook

import json

from core.spreadsheet_finance import CategoryReference, SpreadsheetFinanceService


def make_categories(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Kategori Pengeluaran", "Subkategori"])
    sheet.append(["Makanan", "Makanan", "Restoran", "Cafe", "Makan Siang"])
    sheet.append(["Jual Beli"])
    workbook.save(path)
    return path


def test_category_reference_reads_rows_and_no_subcategory(tmp_path):
    reference = CategoryReference(str(make_categories(tmp_path / "categories.xlsx")))

    assert reference.canonical("makanan", "restoran") == ("Makanan", "Restoran")
    assert reference.canonical("makanan", "makanan") == ("Makanan", "Makanan")
    assert reference.canonical("makanan", "makan siang") == ("Makanan", "Makan Siang")
    assert reference.canonical("Jual Beli", "-") == ("Jual Beli", "-")
    assert reference.canonical("Makanan", "-") is None


def test_export_writes_realbyte_xlsx_and_tsv(tmp_path):
    categories = make_categories(tmp_path / "categories.xlsx")
    service = SpreadsheetFinanceService(
        str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
    )

    result = service.export_batch("2026-07", [{
        "transaction_date": "2026-07-16", "account": "Cash",
        "category": "Makanan", "subcategory": "Restoran", "merchant": "Kedai",
        "amount": 25000, "type": "expense", "description": "Makan",
    }])

    rows = list(load_workbook(result.xlsx_path, read_only=True).active.values)
    assert list(rows[0]) == [
        "Date", "Account", "Category", "Subcategory", "Note", "Amount",
        "Income/Expense", "Description",
    ]
    assert list(rows[1])[:4] == ["16/07/2026", "Cash", "Makanan", "Restoran"]
    tsv = open(result.tsv_path, encoding="utf-8").read()
    assert "16/07/2026\tCash\tMakanan\tRestoran" in tsv
    assert "\tExpense\t" in tsv
    assert result.tsv_path.endswith("2026-07.tsv")
    assert "\r\n" in open(result.tsv_path, "rb").read().decode("utf-8")


def test_history_workbook_is_reported_by_month(tmp_path):
    history = make_categories(tmp_path / "categories.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Account", "Category", "Subcategory", "Note", "Amount", "Income/Expense", "Description"])
    sheet.append(["04/01/2025", "Cash", "Makanan", "Restoran", "Kedai", 10000, "Expense", "Makan"])
    workbook.save(tmp_path / "history.xlsx")
    service = SpreadsheetFinanceService(str(tmp_path / "exports"), str(history), str(tmp_path / "history.xlsx"))

    summary = service.month_summary("2025-04")

    assert summary["expense_total"] == 10000
    assert summary["by_category"] == {"Makanan": 10000}


def test_history_category_examples_prefer_similar_merchant(tmp_path):
    categories = make_categories(tmp_path / "categories.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Account", "Category", "Subcategory", "Note", "Amount", "Income/Expense", "Description"])
    sheet.append(["04/01/2025", "Cash", "Makanan", "Cafe", "-", 16000, "Expense", "Transfer ke SOPUTAN (famima)"])
    sheet.append(["04/02/2025", "Cash", "Makanan", "Grocery", "-", 35000, "Expense", "FamilyMart Kepu Selatan"])
    sheet.append(["04/03/2025", "Cash", "Transportasi", "Bahan Bakar", "Shell", 50000, "Expense", "Bensin"])
    workbook.save(tmp_path / "history.xlsx")
    service = SpreadsheetFinanceService(
        str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
    )

    prompt = service.history_examples_prompt("Family Mart", "cranberry americano")

    assert "Makanan / Cafe" in prompt
    assert "famima" in prompt.casefold() or "FamilyMart" in prompt
    assert "Shell" not in prompt


def test_category_reference_falls_back_to_recommendations_json(tmp_path):
    missing_xlsx = tmp_path / "missing.xlsx"
    recommendations = tmp_path / "category_recommendations.json"
    recommendations.write_text(
        json.dumps({
            "pairs": [["Makanan", "Cafe"], ["Transportasi", "Bahan Bakar"]],
            "phrases": [],
            "keywords": [{
                "match": "soto",
                "category": "Makanan",
                "subcategory": "Restoran",
                "count": 3,
                "confidence": 1.0,
            }],
        }),
        encoding="utf-8",
    )

    reference = CategoryReference(str(missing_xlsx), str(recommendations))

    assert reference.is_loaded()
    assert reference.canonical("makanan", "cafe") == ("Makanan", "Cafe")
    assert reference.canonical("makanan", "restoran") == ("Makanan", "Restoran")
    assert reference.canonical("transportasi", "bahan bakar") == ("Transportasi", "Bahan Bakar")


def test_export_writes_transfer_with_destination_in_category(tmp_path):
    categories = make_categories(tmp_path / "categories.xlsx")
    service = SpreadsheetFinanceService(
        str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
    )

    result = service.export_batch("2026-09", [{
        "transaction_date": "2025-09-02", "account": "Cash",
        "category": "Tabungan - Istri", "subcategory": "-", "merchant": "Nurizka Khoerani",
        "amount": 4002500, "type": "transfer",
        "description": "Transfer dengan BI Fast ke NURIZKA KHOERANI",
    }])

    rows = list(load_workbook(result.xlsx_path, read_only=True).active.values)
    # Destination account in Category → Transfer-Out
    assert list(rows[1]) == [
        "02/09/2025", "Cash", "Tabungan - Istri", "-", "Nurizka Khoerani",
        4002500, "Transfer-Out", "Transfer dengan BI Fast ke NURIZKA KHOERANI",
    ]


def test_export_transfer_with_expense_category_becomes_expenses(tmp_path):
    categories = make_categories(tmp_path / "categories.xlsx")
    # extend categories with transfer expense pair
    from openpyxl import load_workbook, Workbook
    wb = load_workbook(categories)
    wb.active.append(["Penarikan Dana dan Transfer", "Biaya Transfer"])
    wb.save(categories)
    service = SpreadsheetFinanceService(
        str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
    )
    result = service.export_batch("2026-10", [{
        "transaction_date": "2026-07-14", "account": "Cash",
        "category": "Penarikan Dana dan Transfer", "subcategory": "Biaya Transfer",
        "merchant": "Bank Mandiri", "amount": 1502500, "type": "transfer",
        "description": "Transfer BI Fast ke NURIZKA KHOERANI",
    }], force=True)
    rows = list(load_workbook(result.xlsx_path, read_only=True).active.values)
    assert rows[1][6] == "Expense"
    assert rows[1][2] == "Penarikan Dana dan Transfer"
    assert rows[1][3] == "Biaya Transfer"


def test_normalize_finance_account_is_cash_only():
    from core.finance_import import normalize_finance_account, normalize_tx_type

    assert normalize_finance_account("Pribadi") == "Cash"
    assert normalize_finance_account("BCA") == "Cash"
    assert normalize_tx_type("Transfer") == "transfer"
    assert normalize_tx_type("pengeluaran") == "expense"


def test_edit_category_slash_is_kept_without_reresolve(tmp_path):
    from core.db import Database
    from core.finance_import import FinanceImportService
    from core.spreadsheet_finance import SpreadsheetFinanceService

    categories = make_categories(tmp_path / "categories.xlsx")
    db = Database(str(tmp_path / "state.db"))
    batch = db.get_or_create_import_batch("2026-07")
    tx_id = db.stage_transaction(batch["id"], {
        "source_email_id": "E:1",
        "account": "Cash",
        "transaction_date": "2026-07-13",
        "amount": 16000,
        "type": "expense",
        "category": "Makanan",
        "subcategory": "Restoran",
        "merchant": "Warung",
        "description": "makan",
        "needs_review": 0,
    })
    service = FinanceImportService(
        db, None, None,
        spreadsheet_service=SpreadsheetFinanceService(
            str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
        ),
    )
    msg = service.edit_transaction(tx_id, "category", "Makanan / Makan Siang")
    row = db.get_staged_transaction(tx_id)
    assert "Makanan / Makan Siang" in msg
    assert row["category"] == "Makanan"
    assert row["subcategory"] == "Makan Siang"


def test_category_reference_suggests_close_valid_pairs(tmp_path):
    reference = CategoryReference(str(make_categories(tmp_path / "categories.xlsx")))
    suggestions = reference.suggest("Makanan", "Makan Siag", limit=3)
    assert ("Makanan", "Makan Siang") in suggestions
    assert all(reference.canonical(cat, sub) for cat, sub in suggestions)


def test_edit_invalid_category_raises_suggestions(tmp_path):
    from core.db import Database
    from core.finance_import import FinanceImportService, InvalidCategoryError
    from core.spreadsheet_finance import SpreadsheetFinanceService

    categories = make_categories(tmp_path / "categories.xlsx")
    db = Database(str(tmp_path / "state.db"))
    batch = db.get_or_create_import_batch("2026-07")
    tx_id = db.stage_transaction(batch["id"], {
        "source_email_id": "E:2",
        "account": "Cash",
        "transaction_date": "2026-07-13",
        "amount": 16000,
        "type": "expense",
        "category": "Makanan",
        "subcategory": "Restoran",
        "merchant": "Warung",
        "description": "makan",
        "needs_review": 0,
    })
    service = FinanceImportService(
        db, None, None,
        spreadsheet_service=SpreadsheetFinanceService(
            str(tmp_path / "exports"), str(categories), str(tmp_path / "history.xlsx")
        ),
    )
    try:
        service.edit_transaction(tx_id, "category", "Makanan / Tidak Ada")
        assert False, "expected InvalidCategoryError"
    except InvalidCategoryError as exc:
        assert "Tidak Ada" in exc.attempted
        assert exc.suggestions
        assert ("Makanan", "Restoran") in exc.suggestions or all(
            pair[0] == "Makanan" for pair in exc.suggestions
        )

    reply = service.apply_category_pair(tx_id, "Makanan", "Cafe")
    row = db.get_staged_transaction(tx_id)
    assert "Makanan / Cafe" in reply
    assert row["subcategory"] == "Cafe"
