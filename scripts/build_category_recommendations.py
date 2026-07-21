#!/usr/bin/env python3
"""Build category recommendation lookup from Money Manager workbooks.

Reads:
  - Kategori Money Manager.xlsx  (valid category/subcategory pairs)
  - MoneyManager-2025.xlsx       (historical examples)

Writes:
  - config/category_recommendations.json
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.spreadsheet_finance import CategoryReference, _key  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

STOPWORDS = {
    "dengan", "untuk", "dari", "yang", "atau", "pada", "ke", "di", "dan",
    "the", "and", "pembayaran", "transfer", "transaksi", "qris", "qr",
    "via", "uang", "rp", "ke", "oleh",
}

# Extra keywords that rarely appear verbatim in history but map clearly.
SEED_KEYWORDS = [
    ("americano", "Makanan", "Cafe"),
    ("latte", "Makanan", "Cafe"),
    ("espresso", "Makanan", "Cafe"),
    ("cappuccino", "Makanan", "Cafe"),
    ("matcha", "Makanan", "Cafe"),
    ("family mart", "Makanan", "Grocery"),
    ("familymart", "Makanan", "Grocery"),
    ("indosat", "Tagihan dan Langganan", "Seluler"),
    ("telkomsel", "Tagihan dan Langganan", "Seluler"),
    ("xl", "Tagihan dan Langganan", "Seluler"),
    ("rem", "Transportasi", "Servis Kendaraan"),
    ("servis", "Transportasi", "Servis Kendaraan"),
    ("tunjangan", "Gaji", "Tunjangan"),
    ("tukin", "Bonus", "Tukin Tahunan"),
    ("nurizka", "Keluarga", "Istri"),
    ("istri", "Keluarga", "Istri"),
    ("bulanan istri", "Keluarga", "Istri"),
    ("biaya hidup", "Keluarga", "Istri"),
    ("uang belanja", "Keluarga", "Istri"),
    ("bulanan mama", "Keluarga", "Orang Tua"),
    ("mama", "Keluarga", "Orang Tua"),
    ("ratna", "Keluarga", "Orang Tua"),
    ("orang tua", "Keluarga", "Orang Tua"),
]


def build(history_path: Path, categories_path: Path) -> dict:
    reference = CategoryReference(str(categories_path))
    if not reference.is_loaded():
        raise SystemExit(f"Category reference kosong/tidak ditemukan: {categories_path}")

    workbook = load_workbook(history_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(values_only=True))]
    if headers[:8] != [
        "Date", "Account", "Category", "Subcategory", "Note", "Amount",
        "Income/Expense", "Description",
    ]:
        raise SystemExit(f"Header history tidak dikenali: {headers[:8]}")

    phrase_counts: dict[str, Counter] = defaultdict(Counter)
    token_counts: dict[str, Counter] = defaultdict(Counter)
    used_rows = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        pair = reference.canonical(str(row[2] or ""), str(row[3] or "-") or "-")
        if not pair:
            continue
        note = str(row[4] or "").strip()
        description = str(row[7] or "").strip()
        sample = note if note not in {"", "-"} else description
        if not sample or sample == "-":
            continue
        key = _key(sample)
        if len(key) < 4:
            continue
        used_rows += 1
        phrase_counts[key][pair] += 1
        for token in key.split():
            if len(token) >= 4 and token not in STOPWORDS:
                token_counts[token][pair] += 1

    phrases = []
    for key, counter in phrase_counts.items():
        pair, count = counter.most_common(1)[0]
        total = sum(counter.values())
        phrases.append({
            "match": key,
            "category": pair[0],
            "subcategory": pair[1],
            "count": count,
            "confidence": round(count / total, 3),
            "example": key,
        })
    phrases.sort(key=lambda item: (-item["count"], -item["confidence"], item["match"]))

    keywords = []
    for token, counter in token_counts.items():
        pair, count = counter.most_common(1)[0]
        total = sum(counter.values())
        confidence = count / total
        if count < 2 or confidence < 0.6:
            continue
        keywords.append({
            "match": token,
            "category": pair[0],
            "subcategory": pair[1],
            "count": count,
            "total": total,
            "confidence": round(confidence, 3),
        })
    keywords.sort(key=lambda item: (-item["count"], -item["confidence"], item["match"]))

    existing = {item["match"] for item in keywords}
    for match, category, subcategory in SEED_KEYWORDS:
        pair = reference.canonical(category, subcategory)
        if not pair:
            continue
        key = _key(match)
        if key in existing:
            continue
        keywords.append({
            "match": key,
            "category": pair[0],
            "subcategory": pair[1],
            "count": 1,
            "total": 1,
            "confidence": 1.0,
            "seed": True,
        })
        existing.add(key)
    keywords.sort(key=lambda item: (-item["count"], -item["confidence"], item["match"]))

    pairs = sorted(reference.pairs())
    return {
        "version": 1,
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": {
            "history": str(history_path.name),
            "categories": str(categories_path.name),
            "history_rows_used": used_rows,
            "valid_category_pairs": len(pairs),
        },
        "pairs": [[category, subcategory] for category, subcategory in pairs],
        "phrases": phrases,
        "keywords": keywords,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--history",
        default=str(ROOT / "MoneyManager-2025.xlsx"),
        help="Path MoneyManager-2025.xlsx",
    )
    parser.add_argument(
        "--categories",
        default=str(ROOT / "Kategori Money Manager.xlsx"),
        help="Path Kategori Money Manager.xlsx",
    )
    parser.add_argument(
        "--output",
        default=str(ROOT / "config" / "category_recommendations.json"),
        help="Output JSON path",
    )
    args = parser.parse_args()

    payload = build(Path(args.history), Path(args.categories))
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"Wrote {output} "
        f"({len(payload['phrases'])} phrases, {len(payload['keywords'])} keywords, "
        f"{payload['source']['history_rows_used']} history rows)"
    )


if __name__ == "__main__":
    main()
